"""
main.py
PDF OCRツール コアエンジン + CLIエントリーポイント。

GUI（gui.py）と共用するため、処理ロジックを
process_single_pdf() / write_excel() などの関数として公開している。

CLI使用方法:
    python main.py               # inputフォルダを処理
    python main.py path/to/*.pdf # ファイル指定
"""
import os
import sys
import json
import logging
from dataclasses import dataclass, field
from pathlib import Path
from datetime import datetime
from typing import Optional, Callable

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from app_path import APP_DIR
from pdf_reader import extract_text_from_pdf
from ocr import ocr_pdf
from classifier import classify_document
from extractor.invoice import InvoiceExtractor
from extractor.delivery import DeliveryExtractor
from extractor.order import OrderExtractor

# ---------------------------------------------------------------------------
# 定数
# ---------------------------------------------------------------------------

HEADER_COLUMNS = ["ファイル名", "種類", "日付", "金額", "会社名", "番号"]
DETAIL_COLUMNS = ["ファイル名", "行番号", "品名", "数量", "単価", "金額"]

DOC_TYPE_LABELS = {
    "invoice":  "請求書",
    "delivery": "納品書",
    "order":    "発注書",
    "unknown":  "不明",
}

EXTRACTOR_MAP = {
    "invoice":  InvoiceExtractor,
    "delivery": DeliveryExtractor,
    "order":    OrderExtractor,
}

HEADER_BG_COLOR   = "2F5496"
HEADER_FONT_COLOR = "FFFFFF"


# ---------------------------------------------------------------------------
# データクラス
# ---------------------------------------------------------------------------

@dataclass
class ProcessResult:
    """1ファイルの処理結果を保持するデータクラス。"""
    pdf_path:  Path
    status:    str          # "success" | "unknown" | "error"
    doc_type:  str          # "invoice" | "delivery" | "order" | "unknown"
    header:    dict         = field(default_factory=dict)
    details:   list[dict]   = field(default_factory=list)
    error_msg: Optional[str] = None


# ---------------------------------------------------------------------------
# 設定
# ---------------------------------------------------------------------------

def load_config(config_path: str) -> dict:
    """config.jsonを読み込む。"""
    with open(config_path, encoding="utf-8") as f:
        return json.load(f)


def setup_logging(log_dir: str) -> tuple[str, str]:
    """ファイルとコンソールへのログ設定。"""
    os.makedirs(log_dir, exist_ok=True)
    log_file   = os.path.join(log_dir, "log.txt")
    error_file = os.path.join(log_dir, "error_files.txt")

    logging.basicConfig(
        level=logging.INFO,
        format="%(asctime)s [%(levelname)s] %(name)s: %(message)s",
        datefmt="%Y-%m-%d %H:%M:%S",
        handlers=[
            logging.FileHandler(log_file, encoding="utf-8"),
            logging.StreamHandler(sys.stdout),
        ],
    )
    return log_file, error_file


# ---------------------------------------------------------------------------
# エクストラクタ
# ---------------------------------------------------------------------------

def get_extractor(doc_type: str, rules_dir: str):
    """ドキュメント種別に対応するエクストラクタインスタンスを返す。"""
    cls = EXTRACTOR_MAP.get(doc_type)
    if cls is None:
        return None
    rules_file = os.path.join(rules_dir, f"{doc_type}.json")
    if not os.path.exists(rules_file):
        logging.getLogger(__name__).error(f"ルールファイルが見つかりません: {rules_file}")
        return None
    return cls(rules_file)


# ---------------------------------------------------------------------------
# Excel出力
# ---------------------------------------------------------------------------

def _make_header_style(cell, bg: str = HEADER_BG_COLOR, fg: str = HEADER_FONT_COLOR):
    cell.font      = Font(bold=True, color=fg, name="Yu Gothic UI")
    cell.fill      = PatternFill(start_color=bg, end_color=bg, fill_type="solid")
    cell.alignment = Alignment(horizontal="center", vertical="center")


def _auto_fit_columns(ws, min_width: int = 10, max_width: int = 50):
    for col_cells in ws.columns:
        col_letter = get_column_letter(col_cells[0].column)
        max_len = max(
            len(str(cell.value)) if cell.value is not None else 0
            for cell in col_cells
        )
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, min_width), max_width)


def write_excel(headers: list[list], details: list[list], output_path: str) -> None:
    """
    ヘッダ一覧・明細一覧を2シートのExcelファイルに出力する。

    Args:
        headers:     [[ファイル名, 種類, 日付, 金額, 会社名, 番号], ...]
        details:     [[ファイル名, 行番号, 品名, 数量, 単価, 金額], ...]
        output_path: 出力先ファイルパス
    """
    wb = openpyxl.Workbook()

    # --- シート1: ヘッダ ---
    ws1 = wb.active
    ws1.title = "ヘッダ"
    ws1.freeze_panes = "A2"

    for col, title in enumerate(HEADER_COLUMNS, 1):
        _make_header_style(ws1.cell(row=1, column=col, value=title))

    for row_data in headers:
        ws1.append(row_data)

    _auto_fit_columns(ws1)

    # --- シート2: 明細 ---
    ws2 = wb.create_sheet("明細")
    ws2.freeze_panes = "A2"

    for col, title in enumerate(DETAIL_COLUMNS, 1):
        _make_header_style(ws2.cell(row=1, column=col, value=title))

    for row_data in details:
        ws2.append(row_data)

    # 数値書式（数量・単価・金額）
    for row in ws2.iter_rows(min_row=2, max_row=ws2.max_row):
        for col_idx in (4, 5, 6):
            cell = row[col_idx - 1]
            if isinstance(cell.value, (int, float)):
                cell.number_format = '#,##0'

    _auto_fit_columns(ws2)

    os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)
    wb.save(output_path)


# ---------------------------------------------------------------------------
# コア処理（GUI / CLI 共用）
# ---------------------------------------------------------------------------

def process_single_pdf(
    pdf_path: Path,
    rules_dir: str,
    ocr_language: str,
    ocr_dpi: int,
    logger: logging.Logger,
) -> ProcessResult:
    """
    1ファイルを処理して ProcessResult を返す。
    例外は内部で捕捉し status="error" として返すため、呼び出し元は常に結果を受け取れる。
    """
    filename = pdf_path.name

    try:
        # --- テキスト抽出 ---
        text = extract_text_from_pdf(str(pdf_path))
        if not text.strip():
            logger.info(f"テキスト抽出失敗、OCRにフォールバック: {filename}")
            text = ocr_pdf(str(pdf_path), language=ocr_language, dpi=ocr_dpi)

        if not text.strip():
            logger.error(f"テキスト抽出・OCR両方失敗: {filename}")
            return ProcessResult(pdf_path=pdf_path, status="error",
                                 doc_type="unknown", error_msg="テキスト抽出失敗")

        # --- 分類 ---
        doc_type = classify_document(text, rules_dir)
        label    = DOC_TYPE_LABELS.get(doc_type, doc_type)
        logger.info(f"分類: {filename} -> {label}")

        # --- 抽出 ---
        extractor = get_extractor(doc_type, rules_dir)
        if extractor is None:
            return ProcessResult(pdf_path=pdf_path, status="unknown",
                                 doc_type=doc_type, header={}, details=[])

        header  = extractor.extract_header(text)
        details = extractor.extract_details(text)

        logger.info(f"完了: {filename} | 明細{len(details)}行")
        return ProcessResult(
            pdf_path=pdf_path,
            status="success" if doc_type != "unknown" else "unknown",
            doc_type=doc_type,
            header=header,
            details=details,
        )

    except Exception as e:
        logger.error(f"処理エラー: {filename}: {e}", exc_info=True)
        return ProcessResult(pdf_path=pdf_path, status="error",
                             doc_type="unknown", error_msg=str(e))


def results_to_excel_rows(
    results: list[ProcessResult],
) -> tuple[list[list], list[list]]:
    """ProcessResult リストを Excel 出力用の行データに変換する。"""
    headers = []
    details = []

    for r in results:
        h = r.header or {}
        headers.append([
            r.pdf_path.name,
            DOC_TYPE_LABELS.get(r.doc_type, r.doc_type),
            h.get("date"),
            h.get("amount"),
            h.get("company"),
            h.get("number"),
        ])
        for i, d in enumerate(r.details, 1):
            details.append([
                r.pdf_path.name, i,
                d.get("name"), d.get("quantity"),
                d.get("unit_price"), d.get("amount"),
            ])

    return headers, details


# ---------------------------------------------------------------------------
# CLI エントリーポイント
# ---------------------------------------------------------------------------

def main() -> None:
    base_dir    = APP_DIR
    config_path = base_dir / "config.json"
    config      = load_config(str(config_path))

    input_dir  = base_dir / config["input_dir"]
    output_dir = base_dir / config["output_dir"]
    log_dir    = base_dir / config["log_dir"]
    rules_dir  = base_dir / config["rules_dir"]

    input_dir.mkdir(exist_ok=True)
    output_dir.mkdir(exist_ok=True)

    _, error_file = setup_logging(str(log_dir))
    logger = logging.getLogger(__name__)

    logger.info("=" * 60)
    logger.info("PDF OCR Tool 開始 (CLIモード)")
    logger.info(f"入力フォルダ: {input_dir}")
    logger.info("=" * 60)

    # ファイル収集
    if len(sys.argv) > 1:
        pdf_files = [Path(p) for p in sys.argv[1:] if p.lower().endswith(".pdf")]
    else:
        raw = sorted(input_dir.glob("*.pdf")) + sorted(input_dir.glob("*.PDF"))
        seen = set()
        pdf_files = []
        for p in raw:
            if p.name.lower() not in seen:
                seen.add(p.name.lower())
                pdf_files.append(p)

    if not pdf_files:
        logger.warning("処理対象のPDFが見つかりません。")
        return

    logger.info(f"{len(pdf_files)}件のPDFを処理します")

    ocr_language = config.get("ocr_language", "jpn+eng")
    ocr_dpi      = config.get("ocr_dpi", 300)

    results: list[ProcessResult] = []
    for i, pdf_path in enumerate(pdf_files, 1):
        logger.info(f"処理中 ({i}/{len(pdf_files)}): {pdf_path.name}")
        result = process_single_pdf(pdf_path, str(rules_dir), ocr_language, ocr_dpi, logger)
        results.append(result)

    # Excel出力
    output_path = output_dir / config.get("output_file", "result.xlsx")
    headers, details = results_to_excel_rows(results)
    try:
        write_excel(headers, details, str(output_path))
        logger.info(f"Excel出力完了: {output_path}")
    except Exception as e:
        logger.error(f"Excel出力失敗: {e}", exc_info=True)

    # エラーファイルログ
    error_files = [r.pdf_path.name for r in results if r.status == "error"]
    if error_files:
        with open(error_file, "w", encoding="utf-8") as f:
            f.write(f"処理失敗ファイル一覧 ({datetime.now().strftime('%Y-%m-%d %H:%M:%S')})\n")
            f.write("\n".join(error_files))

    success = sum(1 for r in results if r.status == "success")
    unknown = sum(1 for r in results if r.status == "unknown")
    error   = sum(1 for r in results if r.status == "error")
    detail_count = sum(len(r.details) for r in results)

    logger.info("=" * 60)
    logger.info(
        f"処理完了 | 成功:{success}件  不明:{unknown}件  エラー:{error}件  "
        f"明細合計:{detail_count}行"
    )
    logger.info("=" * 60)


if __name__ == "__main__":
    main()
